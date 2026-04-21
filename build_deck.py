import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import io, textwrap
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# ── Palette ──────────────────────────────────────────────────────────────────
NAVY   = RGBColor(0x1B, 0x2A, 0x4A)
RED    = RGBColor(0xC0, 0x39, 0x2B)
AMBER  = RGBColor(0xE6, 0x7E, 0x22)
GREEN  = RGBColor(0x27, 0xAE, 0x60)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LGREY  = RGBColor(0xF4, 0xF6, 0xF7)
MGREY  = RGBColor(0x7F, 0x8C, 0x8D)
DKGREY = RGBColor(0x2C, 0x3E, 0x50)

def rgb(r,g,b): return RGBColor(r,g,b)

# ── Data ─────────────────────────────────────────────────────────────────────
df = pd.read_csv("meridian_snapshot.csv")
panel_end = pd.to_datetime("2026-03-01")
df["report_effective_date"] = pd.to_datetime(df["report_effective_date"])
df["is_active"] = df["report_effective_date"] == panel_end
active   = df[df["is_active"]].copy()
departed = df[~df["is_active"]].copy()
df["tenure_band"] = pd.cut(df["tenure_months"],
    bins=[0,12,24,48,84,999],
    labels=["< 1 yr","1–2 yr","2–4 yr","4–7 yr","7+ yr"])
vol      = df[df["reason_for_departure"] == "Voluntary"]

# ── Helper: embed a matplotlib figure as a picture ───────────────────────────
def fig_to_picture(prs_slide, fig, left, top, width, height):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=180, bbox_inches="tight",
                facecolor=fig.get_facecolor())
    buf.seek(0)
    prs_slide.shapes.add_picture(buf, left, top, width, height)
    plt.close(fig)

# ── Helper: add a text box ────────────────────────────────────────────────────
def add_text(slide, text, left, top, width, height,
             size=14, bold=False, color=DKGREY, align=PP_ALIGN.LEFT,
             wrap=True):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color

# ── Helper: coloured rectangle ───────────────────────────────────────────────
def add_rect(slide, left, top, width, height, fill_rgb, line_rgb=None):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    if line_rgb:
        shape.line.color.rgb = line_rgb
    else:
        shape.line.fill.background()
    return shape

# ── Presentation setup ────────────────────────────────────────────────────────
prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]   # completely blank

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 0 — TITLE
# ════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(blank_layout)
add_rect(sl, Inches(0), Inches(0), Inches(13.33), Inches(7.5), NAVY)
add_rect(sl, Inches(0), Inches(5.2), Inches(13.33), Inches(0.08), GREEN)

add_text(sl, "MERIDIAN TECHNOLOGIES",
         Inches(1), Inches(1.2), Inches(11), Inches(0.8),
         size=18, bold=True, color=rgb(0x27,0xAE,0x60), align=PP_ALIGN.LEFT)

add_text(sl, "People Risk & Workforce Health",
         Inches(1), Inches(2.0), Inches(11), Inches(1.2),
         size=40, bold=True, color=WHITE, align=PP_ALIGN.LEFT)

add_text(sl, "Board Presentation  ·  April 2026  ·  Chief People Officer",
         Inches(1), Inches(3.4), Inches(11), Inches(0.5),
         size=16, bold=False, color=rgb(0xAB,0xB2,0xBD), align=PP_ALIGN.LEFT)

add_text(sl, "CONFIDENTIAL",
         Inches(1), Inches(6.8), Inches(4), Inches(0.4),
         size=10, bold=False, color=rgb(0x7F,0x8C,0x8D), align=PP_ALIGN.LEFT)

add_text(sl, "n = 2,500 employees  ·  36-month panel  ·  Apr 2023 – Mar 2026",
         Inches(7), Inches(6.8), Inches(6), Inches(0.4),
         size=10, bold=False, color=rgb(0x7F,0x8C,0x8D), align=PP_ALIGN.RIGHT)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — STATE OF THE WORKFORCE
# ════════════════════════════════════════════════════════════════════════════
sl1 = prs.slides.add_slide(blank_layout)
add_rect(sl1, Inches(0), Inches(0), Inches(13.33), Inches(7.5), LGREY)
add_rect(sl1, Inches(0), Inches(0), Inches(13.33), Inches(1.1), NAVY)
add_rect(sl1, Inches(0), Inches(1.1), Inches(13.33), Inches(0.06), GREEN)

add_text(sl1, "01  |  STATE OF THE WORKFORCE",
         Inches(0.4), Inches(0.25), Inches(10), Inches(0.6),
         size=22, bold=True, color=WHITE)
add_text(sl1, "April 2026",
         Inches(11.5), Inches(0.32), Inches(1.6), Inches(0.4),
         size=12, color=rgb(0xAB,0xB2,0xBD), align=PP_ALIGN.RIGHT)

# KPI boxes
def kpi_box(slide, left, top, value, label, sub, val_color=NAVY):
    add_rect(slide, left, top, Inches(2.8), Inches(1.6), WHITE)
    add_text(slide, value, left+Inches(0.15), top+Inches(0.1),
             Inches(2.5), Inches(0.8), size=34, bold=True, color=val_color)
    add_text(slide, label, left+Inches(0.15), top+Inches(0.82),
             Inches(2.5), Inches(0.4), size=12, bold=True, color=DKGREY)
    add_text(slide, sub,   left+Inches(0.15), top+Inches(1.2),
             Inches(2.5), Inches(0.35), size=10, color=MGREY)

kpi_box(sl1, Inches(0.35), Inches(1.4),  "1,928",  "Active Employees",   "as of Mar 2026")
kpi_box(sl1, Inches(3.35), Inches(1.4),  "23%",    "36-Month Attrition", "572 departures",     val_color=RED)
kpi_box(sl1, Inches(6.35), Inches(1.4),  "67%",    "Voluntary Exits",    "384 of 572 leavers", val_color=AMBER)
kpi_box(sl1, Inches(9.35), Inches(1.4),  "7.07",   "Avg Engagement",     "scale 1–10")

# Chart 1a: headcount by function
fig1a, ax = plt.subplots(figsize=(5.2, 3.2), facecolor="white")
fn_hc = active["function_name"].value_counts().sort_values()
colors = [str(c) for c in
          ["#1B2A4A" if v == fn_hc.max() else "#AEB6BF" for v in fn_hc]]
ax.barh(fn_hc.index, fn_hc.values, color=colors, edgecolor="white", height=0.65)
for i, v in enumerate(fn_hc.values):
    ax.text(v+5, i, str(v), va="center", fontsize=8.5, color="#2C3E50")
ax.set_xlabel("Active Headcount", fontsize=9)
ax.set_title("Headcount by Function", fontsize=10, fontweight="bold",
             color="#1B2A4A", loc="left", pad=6)
ax.spines[["top","right"]].set_visible(False)
ax.tick_params(labelsize=8.5)
ax.set_facecolor("white")
fig1a.tight_layout(pad=0.5)
fig_to_picture(sl1, fig1a, Inches(0.35), Inches(3.2), Inches(5.5), Inches(3.8))

# Chart 1b: attrition split
fig1b, ax2 = plt.subplots(figsize=(4.0, 3.2), facecolor="white")
labels = ["Voluntary\n(384)", "Involuntary\n(102)", "Silent\n(86)"]
sizes  = [384, 102, 86]
colors2 = ["#C0392B", "#E67E22", "#7F8C8D"]
wedges, texts, autotexts = ax2.pie(
    sizes, labels=labels, colors=colors2, autopct="%1.0f%%",
    startangle=140, pctdistance=0.65,
    textprops={"fontsize": 8.5},
    wedgeprops={"edgecolor":"white","linewidth":2})
for at in autotexts:
    at.set_fontsize(9); at.set_fontweight("bold"); at.set_color("white")
ax2.set_title("Departure Mix (572 total)", fontsize=10,
              fontweight="bold", color="#1B2A4A", loc="left", pad=6)
fig1b.tight_layout(pad=0.5)
fig_to_picture(sl1, fig1b, Inches(6.0), Inches(3.2), Inches(3.8), Inches(3.8))

# Chart 1c: attrition by function bar
fig1c, ax3 = plt.subplots(figsize=(3.5, 3.2), facecolor="white")
func_total = df.groupby("function_name").size()
func_dep   = departed.groupby("function_name").size()
func_attr  = (func_dep/func_total*100).sort_values()
bar_cols = ["#C0392B" if v>24 else "#E67E22" if v>21 else "#27AE60"
            for v in func_attr]
ax3.barh(func_attr.index, func_attr.values, color=bar_cols,
         edgecolor="white", height=0.65)
for i,v in enumerate(func_attr.values):
    ax3.text(v+0.3, i, f"{v:.0f}%", va="center", fontsize=8, color="#2C3E50")
ax3.set_xlabel("Attrition Rate", fontsize=9)
ax3.set_title("Attrition by Function", fontsize=10, fontweight="bold",
              color="#1B2A4A", loc="left", pad=6)
ax3.spines[["top","right"]].set_visible(False)
ax3.tick_params(labelsize=8)
ax3.set_facecolor("white")
fig1c.tight_layout(pad=0.5)
fig_to_picture(sl1, fig1c, Inches(10.0), Inches(3.2), Inches(3.1), Inches(3.8))

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — RISK 1: NEW-HIRE ATTRITION
# ════════════════════════════════════════════════════════════════════════════
sl2 = prs.slides.add_slide(blank_layout)
add_rect(sl2, Inches(0), Inches(0), Inches(13.33), Inches(7.5), LGREY)
add_rect(sl2, Inches(0), Inches(0), Inches(13.33), Inches(1.1), NAVY)
add_rect(sl2, Inches(0), Inches(1.1), Inches(13.33), Inches(0.06), RED)

add_text(sl2, "02  |  RISK 1:  PEOPLE ARE LEAVING BEFORE THEY RAMP",
         Inches(0.4), Inches(0.25), Inches(12), Inches(0.6),
         size=22, bold=True, color=WHITE)

# Headline stat
add_rect(sl2, Inches(0.35), Inches(1.35), Inches(3.8), Inches(2.0), WHITE)
add_text(sl2, "21%", Inches(0.5), Inches(1.42), Inches(3.5), Inches(1.0),
         size=64, bold=True, color=RED)
add_text(sl2, "of new hires leave\nwithin their first year",
         Inches(0.5), Inches(2.35), Inches(3.5), Inches(0.8),
         size=13, bold=False, color=DKGREY)

add_rect(sl2, Inches(0.35), Inches(3.55), Inches(3.8), Inches(1.5), WHITE)
add_text(sl2, "7.8%", Inches(0.5), Inches(3.62), Inches(3.5), Inches(0.7),
         size=38, bold=True, color=GREEN)
add_text(sl2, "attrition for 7+ year\nemployees — 3× lower",
         Inches(0.5), Inches(4.3), Inches(3.5), Inches(0.6),
         size=12, color=DKGREY)

# Chart: tenure band attrition
vr = (vol.groupby("tenure_band", observed=True).size() /
      df.groupby("tenure_band", observed=True).size() * 100).round(1)

fig2, ax = plt.subplots(figsize=(5.5, 3.8), facecolor="white")
bar_colors = ["#C0392B","#E67E22","#E67E22","#F0B27A","#27AE60"]
bars = ax.bar(vr.index, vr.values, color=bar_colors, edgecolor="white", width=0.6)
for b,v in zip(bars, vr.values):
    ax.text(b.get_x()+b.get_width()/2, v+0.3, f"{v}%",
            ha="center", fontsize=11, fontweight="bold", color="#2C3E50")
ax.set_ylim(0,28)
ax.set_ylabel("Voluntary Attrition Rate", fontsize=10)
ax.axhline(vr.mean(), color="#7F8C8D", linestyle="--", linewidth=1.2,
           label=f"Average ({vr.mean():.1f}%)")
ax.legend(fontsize=9, frameon=False)
ax.spines[["top","right"]].set_visible(False)
ax.tick_params(labelsize=10)
ax.set_facecolor("white")
ax.set_title("Voluntary Attrition Rate by Tenure", fontsize=11,
             fontweight="bold", color="#1B2A4A", loc="left", pad=8)
fig2.tight_layout(pad=0.6)
fig_to_picture(sl2, fig2, Inches(4.4), Inches(1.35), Inches(5.5), Inches(4.2))

# Insight bullets
bullets = [
    "Year-one leavers represent a ~$4–8M annual recruiting & lost-productivity cost",
    "Pattern points to an expectation gap at hiring, not a compensation problem",
    "Proposed action: structured 30/60/90 day check-ins + manager accountability",
]
add_rect(sl2, Inches(10.1), Inches(1.35), Inches(3.0), Inches(4.2),
         rgb(0xFF,0xF5,0xF5))
add_text(sl2, "KEY IMPLICATIONS", Inches(10.25), Inches(1.5),
         Inches(2.7), Inches(0.4), size=9, bold=True, color=RED)
for i, b in enumerate(bullets):
    add_text(sl2, f"› {b}", Inches(10.25), Inches(1.95 + i*1.1),
             Inches(2.7), Inches(0.95), size=9.5, color=DKGREY)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — RISK 2: GENDER GAP
# ════════════════════════════════════════════════════════════════════════════
sl3 = prs.slides.add_slide(blank_layout)
add_rect(sl3, Inches(0), Inches(0), Inches(13.33), Inches(7.5), LGREY)
add_rect(sl3, Inches(0), Inches(0), Inches(13.33), Inches(1.1), NAVY)
add_rect(sl3, Inches(0), Inches(1.1), Inches(13.33), Inches(0.06), RED)

add_text(sl3, "03  |  RISK 2:  WOMEN ARE LEAVING MORE AND EARNING LESS",
         Inches(0.4), Inches(0.25), Inches(12), Inches(0.6),
         size=22, bold=True, color=WHITE)

us     = df[df["location"].isin(["SF","NYC","Austin","Remote-US"])]
us_act = active[active["location"].isin(["SF","NYC","Austin","Remote-US"])]
vol_us = us[us["reason_for_departure"]=="Voluntary"]
genders = ["Man","Woman"]
v_rates = [(vol_us[vol_us["gender"]==g].shape[0] /
            us[us["gender"]==g].shape[0]*100) for g in genders]
pay_med = [us_act[us_act["gender"]==g]["base_salary"].median()/1000 for g in genders]

# Attrition gap stat box
add_rect(sl3, Inches(0.35), Inches(1.35), Inches(2.8), Inches(1.7), WHITE)
add_text(sl3, "+3.1pp", Inches(0.45), Inches(1.42),
         Inches(2.6), Inches(0.85), size=46, bold=True, color=RED)
add_text(sl3, "higher voluntary attrition\nfor women vs. men (US)",
         Inches(0.45), Inches(2.2), Inches(2.6), Inches(0.7),
         size=11, color=DKGREY)

add_rect(sl3, Inches(0.35), Inches(3.2), Inches(2.8), Inches(1.7), WHITE)
add_text(sl3, "–$7.2k", Inches(0.45), Inches(3.27),
         Inches(2.6), Inches(0.85), size=42, bold=True, color=RED)
add_text(sl3, "median salary gap,\nwomen vs. men (US)",
         Inches(0.45), Inches(4.05), Inches(2.6), Inches(0.7),
         size=11, color=DKGREY)

add_rect(sl3, Inches(0.35), Inches(5.05), Inches(2.8), Inches(1.2), WHITE)
add_text(sl3, "14.6% → 17.7%",
         Inches(0.45), Inches(5.12), Inches(2.6), Inches(0.55),
         size=17, bold=True, color=AMBER)
add_text(sl3, "Men vs. Women voluntary\nattrition rate (US)",
         Inches(0.45), Inches(5.62), Inches(2.6), Inches(0.55),
         size=10, color=DKGREY)

# Chart: side-by-side attrition + pay
fig3, (ax_l, ax_r) = plt.subplots(1, 2, figsize=(6.4, 3.8), facecolor="white")
cols = ["#2980B9","#C0392B"]

ax_l.bar(genders, v_rates, color=cols, edgecolor="white", width=0.5)
for i,(g,v) in enumerate(zip(genders,v_rates)):
    ax_l.text(i, v+0.2, f"{v:.1f}%", ha="center", fontsize=12,
              fontweight="bold", color="#2C3E50")
ax_l.set_ylim(0,24); ax_l.set_ylabel("Voluntary Attrition Rate", fontsize=9)
ax_l.set_title("Voluntary Attrition (US)", fontsize=10,
               fontweight="bold", color="#1B2A4A", loc="left")
ax_l.spines[["top","right"]].set_visible(False)
ax_l.set_facecolor("white")

ax_r.bar(genders, pay_med, color=cols, edgecolor="white", width=0.5)
for i,(g,v) in enumerate(zip(genders,pay_med)):
    ax_r.text(i, v+1, f"${v:.0f}k", ha="center", fontsize=12,
              fontweight="bold", color="#2C3E50")
ax_r.set_ylim(0,185); ax_r.set_ylabel("Median Base Salary ($k)", fontsize=9)
ax_r.set_title("Median Salary (US)", fontsize=10,
               fontweight="bold", color="#1B2A4A", loc="left")
ax_r.spines[["top","right"]].set_visible(False)
ax_r.set_facecolor("white")

fig3.tight_layout(pad=0.8)
fig_to_picture(sl3, fig3, Inches(3.45), Inches(1.35), Inches(6.0), Inches(4.4))

bullets3 = [
    "Pay gap likely contributes directly to the attrition differential",
    "Fairness risk: visible in Glassdoor data before it reaches the board",
    "Proposed action: comp equity audit (US first) completed by Q3 2026",
]
add_rect(sl3, Inches(9.7), Inches(1.35), Inches(3.3), Inches(4.2),
         rgb(0xFF,0xF5,0xF5))
add_text(sl3, "KEY IMPLICATIONS", Inches(9.85), Inches(1.5),
         Inches(3.0), Inches(0.4), size=9, bold=True, color=RED)
for i, b in enumerate(bullets3):
    add_text(sl3, f"› {b}", Inches(9.85), Inches(1.95 + i*1.1),
             Inches(3.0), Inches(0.95), size=9.5, color=DKGREY)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — BRIGHT SPOTS
# ════════════════════════════════════════════════════════════════════════════
sl4 = prs.slides.add_slide(blank_layout)
add_rect(sl4, Inches(0), Inches(0), Inches(13.33), Inches(7.5), LGREY)
add_rect(sl4, Inches(0), Inches(0), Inches(13.33), Inches(1.1), NAVY)
add_rect(sl4, Inches(0), Inches(1.1), Inches(13.33), Inches(0.06), GREEN)

add_text(sl4, "04  |  WHAT'S WORKING",
         Inches(0.4), Inches(0.25), Inches(12), Inches(0.6),
         size=22, bold=True, color=WHITE)

# 3 bright spot panels
panels = [
    {
        "title": "M&A Integration Succeeded",
        "stat": "79.5%",
        "stat_sub": "of acquired employees\nstill active",
        "body": (
            "The ~2022 acquisition cohort retained at a higher rate than "
            "direct hires (79.5% vs 74.4%) and voluntary attrition is 3pp "
            "below company average. Engagement is on par with the rest of "
            "the workforce. The integration playbook worked."
        ),
        "action": "Playbook to reuse for future M&A",
    },
    {
        "title": "Performance Management Is Functioning",
        "stat": "3.4×",
        "stat_sub": "over-represented\namong leavers (rating 1)",
        "body": (
            "Low performers (rating 1) leave at 3.4× their share of the "
            "active headcount. High performers (rating 4–5) make up 45% of "
            "actives but only 30% of voluntary leavers. The org is naturally "
            "selecting toward talent."
        ),
        "action": "Invest in calibration to protect this signal",
    },
    {
        "title": "Referrals Retain Best",
        "stat": "81.4%",
        "stat_sub": "referral hire\nretention rate",
        "body": (
            "Referral hires retain at 81.4% — 7pp above direct hires "
            "(74.4%) and 5pp above agency hires. With direct hires making "
            "up 1,201 of 2,500 employees, even a modest shift toward "
            "referrals would meaningfully reduce attrition and cost."
        ),
        "action": "Scale referral programme in Eng, Product & CS",
    },
]

for i, p in enumerate(panels):
    lft = Inches(0.35 + i * 4.33)
    add_rect(sl4, lft, Inches(1.35), Inches(4.05), Inches(5.7), WHITE)
    add_rect(sl4, lft, Inches(1.35), Inches(4.05), Inches(0.08), GREEN)

    add_text(sl4, p["title"], lft+Inches(0.2), Inches(1.48),
             Inches(3.7), Inches(0.55), size=12, bold=True, color=NAVY)
    add_text(sl4, p["stat"], lft+Inches(0.2), Inches(2.05),
             Inches(2.2), Inches(0.9), size=42, bold=True, color=GREEN)
    add_text(sl4, p["stat_sub"], lft+Inches(0.2), Inches(2.88),
             Inches(3.7), Inches(0.55), size=10, color=MGREY)
    add_text(sl4, p["body"], lft+Inches(0.2), Inches(3.42),
             Inches(3.7), Inches(2.2), size=10, color=DKGREY)

    add_rect(sl4, lft, Inches(6.62), Inches(4.05), Inches(0.42),
             rgb(0xE8,0xF8,0xF0))
    add_text(sl4, f"› {p['action']}", lft+Inches(0.15), Inches(6.64),
             Inches(3.8), Inches(0.38), size=9.5, bold=True, color=GREEN)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — ASK / NEXT STEPS
# ════════════════════════════════════════════════════════════════════════════
sl5 = prs.slides.add_slide(blank_layout)
add_rect(sl5, Inches(0), Inches(0), Inches(13.33), Inches(7.5), NAVY)
add_rect(sl5, Inches(0), Inches(1.1), Inches(13.33), Inches(0.06), GREEN)

add_text(sl5, "05  |  WHAT WE'RE ASKING FOR",
         Inches(0.4), Inches(0.25), Inches(12), Inches(0.6),
         size=22, bold=True, color=WHITE)

asks = [
    {
        "num": "01",
        "title": "Compensation Equity Review",
        "deadline": "Q3 2026",
        "body": "Mandate a US-wide pay equity audit with external benchmarking. Prioritise women in Engineering and Product where the gap is largest. Board to receive findings before year-end.",
        "color": RED,
    },
    {
        "num": "02",
        "title": "Early-Tenure Intervention Programme",
        "deadline": "Q2 2026",
        "body": "Structured 30/60/90 day manager check-ins for all new hires. Tie manager performance ratings to first-year retention. Pilot in Product and Customer Success first.",
        "color": AMBER,
    },
    {
        "num": "03",
        "title": "Scale the Referral Engine",
        "deadline": "Q2 2026",
        "body": "Referral hires retain 7pp above direct hires. Increase referral bonus budget and set a target of 30% referral hires in Engineering, Product, and Customer Success by year-end.",
        "color": GREEN,
    },
]

for i, a in enumerate(asks):
    top = Inches(1.45 + i * 1.9)
    add_rect(sl5, Inches(0.4), top, Inches(12.5), Inches(1.7),
             rgb(0x22,0x35,0x55))
    add_rect(sl5, Inches(0.4), top, Inches(0.12), Inches(1.7), a["color"])

    add_text(sl5, a["num"], Inches(0.65), top+Inches(0.1),
             Inches(0.6), Inches(0.55), size=26, bold=True, color=a["color"])
    add_text(sl5, a["title"], Inches(1.35), top+Inches(0.15),
             Inches(7.5), Inches(0.5), size=16, bold=True, color=WHITE)
    add_text(sl5, a["body"], Inches(1.35), top+Inches(0.68),
             Inches(9.5), Inches(0.85), size=11, color=rgb(0xAB,0xB2,0xBD))

    add_rect(sl5, Inches(10.9), top+Inches(0.45), Inches(1.8), Inches(0.6),
             a["color"])
    add_text(sl5, a["deadline"], Inches(10.95), top+Inches(0.5),
             Inches(1.7), Inches(0.5), size=13, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER)

add_text(sl5,
    "Data source: Meridian HRIS snapshot (synthetic)  ·  Prepared by CPO Office  ·  April 2026",
    Inches(0.4), Inches(7.1), Inches(12.5), Inches(0.3),
    size=8.5, color=rgb(0x5D,0x6D,0x7E), align=PP_ALIGN.CENTER)

# ── Save ──────────────────────────────────────────────────────────────────────
prs.save("meridian_board_deck.pptx")
print("Saved meridian_board_deck.pptx")
